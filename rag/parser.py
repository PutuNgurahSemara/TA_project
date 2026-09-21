"""
parser.py — Modul parsing dokumen (PDF, DOCX, TXT) menjadi teks bersih dan terstruktur.

Mendukung:
- PDF  : Ekstraksi teks cerdas dengan deteksi tabel berbasis pdfplumber.
         Tabel otomatis dikonversi ke format Markdown Table (| Kolom 1 | Kolom 2 |)
         agar data baris & kolom dipahami secara presisi oleh embedder dan LLM.
- DOCX : Ekstrak paragraf dan tabel via python-docx
- TXT  : Baca langsung, encoding-safe
"""

from pathlib import Path
from typing import Any


def table_to_markdown(table: list[list[Any]]) -> str:
    """
    Konversi representasi list-of-rows menjadi format Markdown Table yang rapi.
    Menghilangkan newline di dalam sel dan menormalkan jumlah kolom.
    """
    if not table or len(table) < 1:
        return ""

    cleaned_rows = []
    for row in table:
        cleaned_row = [" ".join(str(cell).split()) if cell is not None else "" for cell in row]
        if any(cleaned_row):
            cleaned_rows.append(cleaned_row)

    if not cleaned_rows:
        return ""

    max_cols = max(len(r) for r in cleaned_rows)
    if max_cols == 0:
        return ""

    # Pastikan setiap baris memiliki jumlah kolom yang sama
    for r in cleaned_rows:
        while len(r) < max_cols:
            r.append("")

    # Header baris pertama
    headers = [h if h else f"Kolom {j+1}" for j, h in enumerate(cleaned_rows[0])]
    separator = [":---"] * max_cols

    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(separator) + " |",
    ]

    for row in cleaned_rows[1:]:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


def parse_pdf(path: str) -> list[dict[str, Any]]:
    """
    Parse PDF dengan ekstraksi teks cerdas & konversi tabel ke Markdown.
    
    Tabel yang terdeteksi di halaman akan diubah menjadi format Markdown Table,
    sehingga LLM qwen3.5:9b dan embedder bge-m3 dapat memahami relasi
    baris dan kolom secara akurat tanpa kehilangan struktur.
    """
    import pdfplumber

    pages = []
    source_name = Path(path).name

    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            page_content_parts = []

            try:
                tables = page.find_tables()
            except Exception:
                tables = []

            if tables:
                table_bboxes = [t.bbox for t in tables]

                def not_within_tables(obj):
                    v_mid = (obj.get("top", 0) + obj.get("bottom", 0)) / 2
                    h_mid = (obj.get("x0", 0) + obj.get("x1", 0)) / 2
                    return not any(
                        (bbox[0] <= h_mid <= bbox[2]) and (bbox[1] <= v_mid <= bbox[3])
                        for bbox in table_bboxes
                    )

                try:
                    non_table_text = page.filter(not_within_tables).extract_text() or ""
                except Exception:
                    non_table_text = page.extract_text() or ""

                if non_table_text.strip():
                    page_content_parts.append(non_table_text.strip())

                # Ekstrak setiap tabel dan ubah ke Markdown Table
                for t_idx, table_obj in enumerate(tables, start=1):
                    try:
                        table_data = table_obj.extract()
                        md_table = table_to_markdown(table_data)
                        if md_table:
                            page_content_parts.append(f"\n[Tabel {t_idx}]:\n{md_table}")
                    except Exception:
                        pass
            else:
                # Jika tidak ada tabel terdeteksi, ekstrak teks normal
                raw_text = page.extract_text() or ""
                if raw_text.strip():
                    page_content_parts.append(raw_text.strip())

            final_page_text = "\n\n".join(page_content_parts).strip()
            if final_page_text:
                pages.append(
                    {
                        "text": final_page_text,
                        "page_num": i,
                        "source": source_name,
                    }
                )

    return pages


def parse_docx(path: str) -> list[dict[str, Any]]:
    """Parse DOCX (paragraf + tabel) → list of {text, page_num, source}."""
    from docx import Document  # python-docx

    doc = Document(path)
    parts = []

    # Ekstrak paragraf
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())

    # Ekstrak tabel jika ada
    for t_idx, table in enumerate(doc.tables, start=1):
        table_data = []
        for row in table.rows:
            table_data.append([cell.text.strip() for cell in row.cells])
        md_table = table_to_markdown(table_data)
        if md_table:
            parts.append(f"\n[Tabel {t_idx}]:\n{md_table}")

    full_text = "\n\n".join(parts)
    if not full_text:
        return []

    return [{"text": full_text, "page_num": 1, "source": Path(path).name}]


def parse_txt(path: str) -> list[dict[str, Any]]:
    """Parse TXT → list of {text, page_num, source}."""
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read().strip()

    if not text:
        return []

    return [{"text": text, "page_num": 1, "source": Path(path).name}]


def parse_document(path: str) -> list[dict[str, Any]]:
    """Router: pilih parser berdasarkan ekstensi file."""
    ext = Path(path).suffix.lower()
    parsers = {
        ".pdf": parse_pdf,
        ".docx": parse_docx,
        ".txt": parse_txt,
    }
    if ext not in parsers:
        raise ValueError(
            f"Format file tidak didukung: '{ext}'. "
            f"Format yang diterima: {', '.join(parsers.keys())}"
        )
    return parsers[ext](path)
